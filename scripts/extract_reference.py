#!/usr/bin/env python3
"""
Extraction de la référence budgétaire (Prévisionnel) depuis le classeur
Excel (.xlsx) ou OpenDocument (.ods) de suivi analytique de devenir·art,
vers un fichier reference-budget.json exploitable par l'application web.

Logique (3 niveaux : poste -> groupe (optionnel) -> catégorie) :
- Un onglet par axe (Fonctionnement + 6 actions).
- Dans chaque onglet, la colonne A porte les libellés, la colonne C le
  Prévisionnel, la colonne D le Réalisé.
- Une ligne "poste comptable" commence par un nombre à deux chiffres suivi
  d'un tiret (ex. "60 – Achats") : elle sert uniquement de séparateur, ce
  n'est jamais une catégorie de saisie.
- Une ligne "catégorie feuille" (ex. "carburant", "Fournitures bureau") est
  identifiable par le fait que sa cellule Prévisionnel OU Réalisé contient une
  formule SUM() sur une PLAGE contiguë (ex. =SUM(C11:C18), ou son équivalent
  ODF =SUM([.C11:.C18])) : cette plage référence les lignes de transaction
  individuelles juste en dessous.
- Une ligne intermédiaire de regroupement ("groupe", ex. "Achats matières &
  fournitures", "Déplacement, voyages") agrège plusieurs catégories : sa
  formule additionne des cellules isolées (=SUM(C10+C19+C36) ou
  =C10+C19+C36, sans plage) au lieu d'une plage contiguë. Ce n'est pas une
  catégorie de saisie, mais son libellé est conservé et associé à chacune
  des catégories qu'elle référence (le champ "groupe" du JSON de sortie),
  pour que l'app puisse les afficher regroupées comme dans le classeur.
  Une catégorie qui n'est référencée par aucun groupe est un enfant direct
  du poste (groupe = null).
- Exception : le poste "63 – Impôts & taxes" saisit ses 2 catégories en
  valeur brute, sans formule SUM du tout (pas de détail de transactions en
  dessous) — on les traite comme des feuilles quand même, sans groupe.
- On ne lit jamais les lignes de transaction individuelles elles-mêmes (les
  valeurs brutes en dessous d'une catégorie) : seul le sous-total de la
  catégorie (son Prévisionnel figé pour l'année) nous intéresse.

Recettes (comptes 70 à 79) : si le classeur contient un onglet nommé
"produits" (recherche insensible à la casse/aux espaces), ses 7 postes
officiels sont extraits sous la clé "Recettes" du JSON de sortie. Cet
onglet n'a pas (encore) la même rigueur de structure que les onglets de
charges (pas de formule SUM(plage) fiable sur les sous-catégories) : seul
le niveau poste est extrait automatiquement, le nom du poste servant de
catégorie par défaut — les sous-catégories plus fines (ex. « Subvention
DRAC ») sont à ajouter à la main dans l'app.

Le script ne modifie jamais le classeur source (ouverture en lecture seule).

Usage :
    py scripts/extract_reference.py "chemin/vers/BP_....xlsx" [--annee 2026] [--out public/reference-budget.json]
    py scripts/extract_reference.py "chemin/vers/BP_....ods"  [--annee 2026] [--out public/reference-budget.json]

Ré-exécutable à tout moment sur un nouveau classeur (même structure de
gabarit, nouveaux montants) : il suffit de repasser le chemin du nouveau
fichier, Excel ou OpenDocument (LibreOffice/OnlyOffice).
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

# Ordre et libellés des 7 axes tels qu'ils doivent apparaitre dans l'app,
# associés au nom exact de l'onglet dans le classeur source.
AXES: list[tuple[str, str]] = [
    ("Fonctionnement", "Fonctionnement"),
    ("Commissions", "Commissions"),
    ("Points Infos", "Points Infos"),
    ("RRAV", "RRAV"),
    ("Coopération et Visibilité", "Coopération et Visibilité"),
    ("Navettes de l'art", "navettes de l'art"),
    ("Représentation", "Représentation"),
]

POSTE_RE = re.compile(r"^\s*\d{2}\s*[–‒\-]\s*.+")

# Formule SUM() sur une plage contiguë -> catégorie feuille.
# Variante Excel :        =SUM(C11:C18)
# Variante ODF (LibreOffice/OnlyOffice) : of:=SUM([.C11:.C18])
LEAF_SUM_RANGE_RE_XLSX = re.compile(r"^=SUM\([A-Za-z]+\d+:[A-Za-z]+\d+\)\s*$")
LEAF_SUM_RANGE_RE_ODS = re.compile(r"^of:=SUM\(\[\.[A-Za-z]+\d+:\.[A-Za-z]+\d+\]\)\s*$")

# Référence de cellule isolée dans une formule de regroupement, ex. le "C10",
# "C19" et "C36" de "=SUM(C10+C19+C36)" ou "of:=SUM([.C10]+[.C19]+[.C36])" —
# la lettre de colonne n'importe pas, seul le numéro de ligne compte.
CELL_REF_RE = re.compile(r"[A-Za-z]+(\d+)")

STOP_LABELS = {
    "TOTAL DES CHARGES",
    "TOTAL DES PRODUITS",
    "PART BUDGET GLOBAL",
    "TOTAL GENERAL",
    "TOTAL GÉNÉRAL",
    "CONTRIBUTIONS VOLONTAIRES EN NATURE",
}

POSTE_PREFIX_RE = re.compile(r"^\s*\d+\s*[–‒\-]\s*")


def nom_sans_prefixe_poste(label: str) -> str:
    """« 70 – Ventes » -> « Ventes » : sert de nom de catégorie par défaut
    pour les recettes, dont le classeur ne détaille pas encore les
    sous-catégories aussi finement que pour les charges."""
    return POSTE_PREFIX_RE.sub("", label).strip()


def trouver_feuille_produits(sheetnames: list[str]) -> str | None:
    for name in sheetnames:
        if name.strip().lower() == "produits":
            return name
    return None

# Le poste "63 – Impôts & taxes" ne suit pas le schéma habituel : ses 2
# catégories ("Impôts & taxes sur rémunération", "Autres impôts & taxes")
# sont saisies directement en valeur brute, sans ligne de sous-total SUM(plage)
# au-dessus d'elles (il n'y a jamais de détail de transactions en dessous).
# On les traite donc comme des feuilles même sans formule SUM.
FLAT_LEAF_POSTE_RE = re.compile(r"^\s*63\s*[–‒\-]")


def clean_label(value) -> str:
    if value is None:
        return ""
    return str(value).replace("\xa0", " ").strip()


def is_leaf_category(formula_value) -> bool:
    if not isinstance(formula_value, str):
        return False
    v = formula_value.strip()
    return bool(LEAF_SUM_RANGE_RE_XLSX.match(v)) or bool(LEAF_SUM_RANGE_RE_ODS.match(v))


def est_feuille(poste_courant: str, c_formula, d_formula) -> bool:
    if is_leaf_category(c_formula) or is_leaf_category(d_formula):
        return True
    return bool(FLAT_LEAF_POSTE_RE.match(poste_courant))


def lignes_referencees(formula_value) -> list[int]:
    """Numéros de ligne référencés par une formule de regroupement (ex. les
    3 lignes de =SUM(C10+C19+C36) -> [10, 19, 36]). Vide si pas de formule."""
    if not isinstance(formula_value, str):
        return []
    return [int(n) for n in CELL_REF_RE.findall(formula_value)]


class SuiviGroupes:
    """Mémorise, pendant le parcours séquentiel d'un onglet, quelle ligne de
    regroupement ("groupe") est la mère de quelle ligne de catégorie —
    construit au fil de l'eau à partir des formules des lignes de
    regroupement (une catégorie n'a jamais un numéro de ligne inférieur à
    son groupe, donc un simple parcours du haut vers le bas suffit)."""

    def __init__(self) -> None:
        self._groupe_de: dict[int, str] = {}

    def enregistrer_groupe(self, label: str, c_formula, d_formula) -> None:
        for n in {*lignes_referencees(c_formula), *lignes_referencees(d_formula)}:
            self._groupe_de[n] = label

    def groupe_de(self, row_num: int) -> str | None:
        return self._groupe_de.get(row_num)


# ---------------------------------------------------------------------------
# Lecteur .xlsx (openpyxl)
# ---------------------------------------------------------------------------

def _charger_classeur_xlsx(path: Path):
    try:
        import openpyxl
    except ImportError:
        print(
            "Le module 'openpyxl' est requis pour lire un .xlsx. Installez-le avec : py -m pip install openpyxl",
            file=sys.stderr,
        )
        sys.exit(1)

    # data_only=False : on veut les formules pour détecter les catégories feuilles.
    wb_formulas = openpyxl.load_workbook(path, data_only=False, read_only=True)
    # data_only=True : on veut les valeurs mises en cache par Excel pour le Prévisionnel.
    wb_values = openpyxl.load_workbook(path, data_only=True, read_only=True)

    def extraire(sheet_name: str) -> list[dict]:
        ws_formulas = wb_formulas[sheet_name]
        ws_values = wb_values[sheet_name]
        lignes: list[dict] = []
        poste_courant: str | None = None
        groupes = SuiviGroupes()

        for row in range(1, ws_formulas.max_row + 1):
            label = clean_label(ws_formulas.cell(row=row, column=1).value)
            if not label:
                continue
            if label.upper() in STOP_LABELS:
                break
            if POSTE_RE.match(label):
                poste_courant = label
                continue
            if poste_courant is None:
                continue

            c_formula = ws_formulas.cell(row=row, column=3).value
            d_formula = ws_formulas.cell(row=row, column=4).value

            if est_feuille(poste_courant, c_formula, d_formula):
                prevu_value = ws_values.cell(row=row, column=3).value
                prevu = round(float(prevu_value), 2) if isinstance(prevu_value, (int, float)) else 0.0
                lignes.append({
                    "poste": poste_courant,
                    "groupe": groupes.groupe_de(row),
                    "categorie": label,
                    "prevu": prevu,
                })
            else:
                groupes.enregistrer_groupe(label, c_formula, d_formula)

        return lignes

    def extraire_produits(sheet_name: str) -> list[dict]:
        # Feuille "produits" : structure différente des feuilles de charges
        # (colonne A = libellé, colonne B = Prévisionnel — pas C). Seul le
        # niveau poste (70 à 79) est fiable pour l'instant (formules pas
        # encore posées sur les sous-catégories dans le classeur) ; le nom du
        # poste sert de catégorie par défaut, à affiner à la main dans l'app.
        ws_values = wb_values[sheet_name]
        lignes: list[dict] = []
        for row in range(1, ws_values.max_row + 1):
            label = clean_label(ws_values.cell(row=row, column=1).value)
            if not label:
                continue
            if label.upper() in STOP_LABELS:
                break
            if POSTE_RE.match(label):
                prevu_value = ws_values.cell(row=row, column=2).value
                prevu = round(float(prevu_value), 2) if isinstance(prevu_value, (int, float)) else 0.0
                lignes.append({
                    "poste": label,
                    "groupe": None,
                    "categorie": nom_sans_prefixe_poste(label),
                    "prevu": prevu,
                })
        return lignes

    return wb_formulas.sheetnames, extraire, extraire_produits


# ---------------------------------------------------------------------------
# Lecteur .ods (odfpy)
# ---------------------------------------------------------------------------

def _charger_classeur_ods(path: Path):
    try:
        from odf.opendocument import load
        from odf.table import Table, TableRow, TableCell
        from odf import teletype
    except ImportError:
        print(
            "Le module 'odfpy' est requis pour lire un .ods. Installez-le avec : py -m pip install odfpy",
            file=sys.stderr,
        )
        sys.exit(1)

    doc = load(str(path))
    tables = doc.spreadsheet.getElementsByType(Table)
    tables_par_nom = {t.getAttribute("name"): t for t in tables}

    def cellules_de_la_ligne(row, max_colonnes: int = 10) -> list:
        # Une cellule ODF peut représenter N colonnes identiques d'affilée
        # (table:number-columns-repeated), typiquement les cellules vides de
        # remplissage entre deux valeurs saisies : il faut les déplier pour
        # que la position (A=1, C=3, D=4...) reste correcte.
        plat: list = []
        for cell in row.getElementsByType(TableCell):
            repet = int(cell.getAttribute("numbercolumnsrepeated") or 1)
            plat.extend([cell] * min(repet, max_colonnes))
            if len(plat) >= max_colonnes:
                break
        return plat

    def extraire(sheet_name: str) -> list[dict]:
        table = tables_par_nom[sheet_name]
        lignes: list[dict] = []
        poste_courant: str | None = None
        groupes = SuiviGroupes()

        # Le numéro de ligne réel doit être reconstitué car table:
        # number-rows-repeated compresse les longues plages de lignes vides
        # (surtout en fin de feuille) en un seul élément XML : il faut
        # avancer le compteur de son "repeat" pour que les références de
        # cellules dans les formules (ex. C10) continuent de correspondre au
        # bon numéro de ligne pour les lignes de contenu qui suivent.
        row_num = 0
        for row in table.getElementsByType(TableRow):
            repeat = int(row.getAttribute("numberrowsrepeated") or 1)
            row_num += repeat
            if repeat > 1:
                continue  # bloc de lignes vides répétées : jamais de libellé

            cellules = cellules_de_la_ligne(row)
            if not cellules:
                continue
            label = clean_label(teletype.extractText(cellules[0]))
            if not label:
                continue
            if label.upper() in STOP_LABELS:
                break
            if POSTE_RE.match(label):
                poste_courant = label
                continue
            if poste_courant is None:
                continue

            c_cell = cellules[2] if len(cellules) > 2 else None
            d_cell = cellules[3] if len(cellules) > 3 else None
            c_formula = c_cell.getAttribute("formula") if c_cell is not None else None
            d_formula = d_cell.getAttribute("formula") if d_cell is not None else None

            if est_feuille(poste_courant, c_formula, d_formula):
                c_value = c_cell.getAttribute("value") if c_cell is not None else None
                prevu = round(float(c_value), 2) if c_value not in (None, "") else 0.0
                lignes.append({
                    "poste": poste_courant,
                    "groupe": groupes.groupe_de(row_num),
                    "categorie": label,
                    "prevu": prevu,
                })
            else:
                groupes.enregistrer_groupe(label, c_formula, d_formula)

        return lignes

    def extraire_produits(sheet_name: str) -> list[dict]:
        table = tables_par_nom[sheet_name]
        lignes: list[dict] = []
        for row in table.getElementsByType(TableRow):
            repeat = int(row.getAttribute("numberrowsrepeated") or 1)
            if repeat > 1:
                continue
            cellules = cellules_de_la_ligne(row)
            if not cellules:
                continue
            label = clean_label(teletype.extractText(cellules[0]))
            if not label:
                continue
            if label.upper() in STOP_LABELS:
                break
            if POSTE_RE.match(label):
                b_cell = cellules[1] if len(cellules) > 1 else None
                b_value = b_cell.getAttribute("value") if b_cell is not None else None
                prevu = round(float(b_value), 2) if b_value not in (None, "") else 0.0
                lignes.append({
                    "poste": label,
                    "groupe": None,
                    "categorie": nom_sans_prefixe_poste(label),
                    "prevu": prevu,
                })
        return lignes

    return list(tables_par_nom.keys()), extraire, extraire_produits


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("classeur_path", help="Chemin vers le fichier source (.xlsx ou .ods)")
    parser.add_argument("--annee", type=int, default=None, help="Année budgétaire (déduite du nom de fichier si absent)")
    parser.add_argument("--out", default="public/reference-budget.json", help="Fichier JSON de sortie")
    args = parser.parse_args()

    classeur_path = Path(args.classeur_path)
    if not classeur_path.exists():
        print(f"Fichier introuvable : {classeur_path}", file=sys.stderr)
        sys.exit(1)

    annee = args.annee
    if annee is None:
        m = re.search(r"20\d{2}", classeur_path.stem)
        annee = int(m.group(0)) if m else None
    if annee is None:
        print("Impossible de déduire l'année depuis le nom de fichier : précisez --annee.", file=sys.stderr)
        sys.exit(1)

    suffixe = classeur_path.suffix.lower()
    if suffixe == ".xlsx":
        print(f"Lecture (lecture seule, Excel) de {classeur_path} ...")
        sheetnames, extraire, extraire_produits = _charger_classeur_xlsx(classeur_path)
    elif suffixe == ".ods":
        print(f"Lecture (lecture seule, OpenDocument) de {classeur_path} ...")
        sheetnames, extraire, extraire_produits = _charger_classeur_ods(classeur_path)
    else:
        print(f"Format non pris en charge : {suffixe} (attendu : .xlsx ou .ods)", file=sys.stderr)
        sys.exit(1)

    result: dict = {"annee": annee}
    total_categories = 0

    for axe_label, sheet_name in AXES:
        if sheet_name not in sheetnames:
            print(f"  ! Onglet '{sheet_name}' introuvable, ignoré.", file=sys.stderr)
            continue
        lignes = extraire(sheet_name)
        result[axe_label] = lignes
        total_categories += len(lignes)
        print(f"  - {axe_label}: {len(lignes)} catégories, prévu total = {sum(l['prevu'] for l in lignes):.2f} EUR")

    feuille_produits = trouver_feuille_produits(sheetnames)
    if feuille_produits:
        lignes_produits = extraire_produits(feuille_produits)
        result["Recettes"] = lignes_produits
        total_categories += len(lignes_produits)
        print(
            f"  - Recettes: {len(lignes_produits)} postes, prévu total = "
            f"{sum(l['prevu'] for l in lignes_produits):.2f} EUR"
        )
    else:
        print("  (pas d'onglet 'produits' dans ce fichier — aucune recette extraite)")

    out_path = Path(args.out)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"\n{total_categories} catégories extraites -> {out_path}")


if __name__ == "__main__":
    main()
